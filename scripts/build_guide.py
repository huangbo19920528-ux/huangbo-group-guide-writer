from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


CHINESE_FONT = "楷体"
LATIN_FONT = "Times New Roman"
BLACK = "000000"
GRAY = "7F7F7F"
LIGHT_GRAY = "E7E6E6"
GRID = "D9D9D9"
NOTICE = "黄博课题组内部资料 请勿外传"


def set_run_font(run, size=12, bold=False, italic=False, color=BLACK):
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), CHINESE_FONT)
    rfonts.set(qn("w:cs"), LATIN_FONT)
    return run


def configure_style(style, size, *, bold=False, before=0, after=4,
                    line_spacing=1.35, first_line=0, keep_next=False):
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), CHINESE_FONT)
    rfonts.set(qn("w:cs"), LATIN_FONT)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing
    pf.first_line_indent = Pt(first_line)
    pf.keep_with_next = keep_next


def remove_title_border(style):
    ppr = style.element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    set_run_font(run, size=9, color=GRAY)
    run._r.extend([begin, text, separate, value, end])


def get_or_add_style(styles, name):
    return styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shading = tcpr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tcpr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_borders(table, color=GRID, size="6"):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    trpr.append(marker)


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def set_cell_text(cell, text, *, size=10.2, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.2
    set_run_font(paragraph.add_run(str(text)), size=size, bold=bold)


def add_text_paragraph(doc, text, style_name, *, indent=False, size=12,
                       color=BLACK, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph = doc.add_paragraph(style=style_name)
    paragraph.alignment = align
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)
    set_run_font(paragraph.add_run(text), size=size, color=color)
    return paragraph


def add_table(doc, block):
    headers = block.get("headers")
    rows = block.get("rows")
    if not headers or not isinstance(rows, list):
        raise ValueError("table requires non-empty headers and rows")
    columns = len(headers)
    if any(len(row) != columns for row in rows):
        raise ValueError("every table row must have the same number of cells as headers")
    widths = block.get("widths_cm") or [15.7 / columns] * columns
    if len(widths) != columns or sum(widths) > 16.0:
        raise ValueError("table widths_cm must match columns and total no more than 16.0 cm")
    size = float(block.get("font_size", 10.2))
    center_columns = set(block.get("center_columns", [0]))

    table = doc.add_table(rows=1, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        shade_cell(cell, LIGHT_GRAY)
        set_cell_text(cell, header, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.width = Cm(widths[index])
            alignment = WD_ALIGN_PARAGRAPH.CENTER if index in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cell, value, size=size, align=alignment)

    set_table_borders(table)
    for row in table.rows:
        prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def build(spec_path, output_path):
    spec_path = Path(spec_path).resolve()
    output_path = Path(output_path).resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    title_text = str(spec.get("title", "")).strip()
    if not title_text:
        raise ValueError("title is required")
    blocks = spec.get("blocks")
    if not isinstance(blocks, list):
        raise ValueError("blocks must be a list")
    if any(NOTICE in json.dumps(block, ensure_ascii=False) for block in blocks):
        raise ValueError("the internal notice is inserted automatically; remove it from blocks")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.footer_distance = Cm(0.8)

    doc.core_properties.title = title_text
    doc.core_properties.subject = str(spec.get("subject", spec.get("subtitle", "")))
    doc.core_properties.author = "黄博课题组"

    styles = doc.styles
    configure_style(styles["Normal"], 12)
    configure_style(styles["Title"], 22, bold=True, after=8, line_spacing=1.15)
    remove_title_border(styles["Title"])
    configure_style(styles["Heading 1"], 16, bold=True, before=9, after=5, line_spacing=1.15, keep_next=True)
    configure_style(styles["Heading 2"], 14, bold=True, before=7, after=4, line_spacing=1.15, keep_next=True)
    configure_style(styles["Heading 3"], 12, bold=True, before=5, after=3, line_spacing=1.15, keep_next=True)
    configure_style(get_or_add_style(styles, "正文缩进"), 12, first_line=24)
    configure_style(get_or_add_style(styles, "正文"), 12)
    configure_style(get_or_add_style(styles, "列表正文"), 12, after=2, line_spacing=1.3)
    configure_style(get_or_add_style(styles, "说明"), 10, after=3, line_spacing=1.2)
    configure_style(get_or_add_style(styles, "示例正文"), 10.5, after=4, line_spacing=1.25)

    section.header.paragraphs[0].text = ""
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("第 "), size=9, color=GRAY)
    add_field(footer, " PAGE ")
    set_run_font(footer.add_run(" 页"), size=9, color=GRAY)

    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.paragraph_format.space_before = Pt(16)
    notice.paragraph_format.space_after = Pt(18)
    set_run_font(notice.add_run(NOTICE), size=14, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run(title_text), size=22, bold=True)
    subtitle_text = str(spec.get("subtitle", "")).strip()
    if subtitle_text:
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(18)
        set_run_font(subtitle.add_run(subtitle_text), size=15)

    for block in blocks:
        kind = block.get("type")
        if kind in {"h1", "h2", "h3"}:
            level = int(kind[-1])
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            set_run_font(paragraph.add_run(str(block["text"])), size={1: 16, 2: 14, 3: 12}[level], bold=True)
        elif kind == "p":
            indent = bool(block.get("indent", True))
            add_text_paragraph(doc, str(block["text"]), "正文缩进" if indent else "正文", indent=indent)
        elif kind == "note":
            add_text_paragraph(doc, str(block["text"]), "说明", size=10, color=GRAY, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif kind == "example":
            paragraph = add_text_paragraph(doc, str(block["text"]), "示例正文", size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.paragraph_format.right_indent = Cm(0.5)
        elif kind in {"bullets", "numbers"}:
            for index, item in enumerate(block.get("items", []), 1):
                prefix = "• " if kind == "bullets" else f"{index}. "
                paragraph = add_text_paragraph(doc, prefix + str(item), "列表正文", align=WD_ALIGN_PARAGRAPH.LEFT)
                paragraph.paragraph_format.left_indent = Cm(0.65)
                paragraph.paragraph_format.first_line_indent = Cm(-0.45)
        elif kind == "table":
            add_table(doc, block)
        elif kind == "image":
            image_path = Path(str(block["path"]))
            if not image_path.is_absolute():
                image_path = spec_path.parent / image_path
            if not image_path.exists():
                raise FileNotFoundError(image_path)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = bool(block.get("caption") or block.get("description"))
            paragraph.add_run().add_picture(str(image_path), width=Cm(float(block.get("width_cm", 15.0))))
            caption_text = str(block.get("caption", "")).strip()
            if caption_text:
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.keep_with_next = bool(block.get("description"))
                set_run_font(caption.add_run(caption_text), size=10)
            description = str(block.get("description", "")).strip()
            if description:
                add_text_paragraph(doc, description, "正文缩进", indent=True)
        elif kind == "page_break":
            doc.add_page_break()
        else:
            raise ValueError(f"unsupported block type: {kind}")

    version = str(spec.get("version", "")).strip()
    if version:
        add_text_paragraph(doc, f"版本日期 {version}", "说明", size=10, color=GRAY, align=WD_ALIGN_PARAGRAPH.LEFT)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(output_path)


def main():
    parser = argparse.ArgumentParser(description="Build a Huang Bo Research Group guide DOCX from JSON")
    parser.add_argument("spec")
    parser.add_argument("output")
    args = parser.parse_args()
    build(args.spec, args.output)


if __name__ == "__main__":
    main()
